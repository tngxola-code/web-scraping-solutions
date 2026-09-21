"""Generated test fixtures (tiny xlsx / csv / pdf).

Fixtures are generated in code so the repository never ships binary blobs.
The PDF fixture uses reportlab when available and otherwise falls back to a
hand-built minimal (but valid) PDF document.
"""

from __future__ import annotations

import os
from typing import Optional

FIXTURE_DIR = os.path.dirname(__file__)


# ----------------------------------------------------------------------
# Spreadsheet fixtures
# ----------------------------------------------------------------------
def make_xlsx(path: str, multi_sheet: bool = True, with_formula: bool = True) -> str:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Clauses"
    ws.append(["Clause No", "Clause Text", "Notes"])
    ws.append(["1", "Parties must comply with  the agreement.", "first"])
    ws.append(["2", "A clause that spans\nmultiple lines stays together.", None])
    ws.append(["3", "  Leading and trailing whitespace  ", "dup"])
    ws.append(["3", "  Leading and trailing whitespace  ", "dup"])
    if with_formula:
        ws.append(["4", "Total rows:", "=COUNTA(A2:A5)"])
    if multi_sheet:
        ws2 = wb.create_sheet("Extra")
        ws2.append(["Ref", "Value"])
        ws2.append(["A", "alpha"])
        ws2.append(["B", "beta"])
    wb.save(path)
    return path


def make_empty_xlsx(path: str) -> str:
    from openpyxl import Workbook

    wb = Workbook()
    wb.active.title = "Empty"
    wb.save(path)
    return path


def make_csv(path: str) -> str:
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("Clause No,Clause Text\n")
        fh.write("1,First clause text\n")
        fh.write("2,Second clause text\n")
        fh.write('3,"Third clause, with comma"\n')
    return path


def make_empty_csv(path: str) -> str:
    with open(path, "w", encoding="utf-8", newline="") as fh:
        fh.write("Clause No,Clause Text\n")
    return path


# ----------------------------------------------------------------------
# PDF fixtures
# ----------------------------------------------------------------------
_PAGE_HEADER = "ABCC Agreement Clauses"
_PAGE_FOOTER = "Australian Building and Construction Commission"
_CLAUSE_LINES = [
    "1. This is the first agreement clause and it continues",
    "onto a second physical line of the same clause.",
    "2. This is the second agreement clause.",
]

# lines: (page1 lines, page2 lines) — clause 3 starts at the bottom of page 1
# and continues on page 2 to exercise cross-page continuation (PDF-009).
PAGE1_LINES = (
    [_PAGE_HEADER]
    + _CLAUSE_LINES
    + ["3. A long clause beginning on page one"]
    + [_PAGE_FOOTER, "Page 1"]
)
PAGE2_LINES = [
    _PAGE_HEADER,
    "and finishing on page two.",
    "4. Final clause.",
    _PAGE_FOOTER,
    "Page 2",
]


def _minimal_pdf_lines(lines: list[str]) -> bytes:
    """Build a minimal valid one-page PDF containing the given text lines."""

    def esc(s: str) -> str:
        return s.replace("\\", r"\\").replace("(", r"\(").replace(")", r"\)")

    ops = ["BT", "/F1 10 Tf", "14 TL", "72 740 Td"]
    for i, line in enumerate(lines):
        if i:
            ops.append("T*")
        ops.append(f"({esc(line)}) Tj")
    ops.append("ET")
    stream = "\n".join(ops).encode("latin-1", "replace")

    objects: list[bytes] = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Courier >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i + obj + b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref_pos,
    )
    return bytes(out)


def make_pdf(path: str) -> str:
    """Two-page PDF with repeated header/footer and a cross-page clause."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        c = canvas.Canvas(path, pagesize=letter)
        for lines in (PAGE1_LINES, PAGE2_LINES):
            y = 740
            for line in lines:
                c.drawString(72, y, line)
                y -= 14
            c.showPage()
        c.save()
        return path
    except ImportError:
        # fallback: single-page minimal valid PDF bytes (no reportlab)
        with open(path, "wb") as fh:
            fh.write(_minimal_pdf_lines(_CLAUSE_LINES))
        return path


def make_scanned_pdf(path: str) -> str:
    """A valid PDF page with no embedded text (for PDF-005 behaviour)."""
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.pdfgen import canvas

        c = canvas.Canvas(path, pagesize=letter)
        c.rect(72, 700, 200, 40)  # drawing only, no text operators
        c.showPage()
        c.save()
        return path
    except ImportError:
        with open(path, "wb") as fh:
            fh.write(_minimal_pdf_lines([]))
        return path
