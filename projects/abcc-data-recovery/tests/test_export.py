"""Export tests (SPEC section 22): workbook creation, expected sheets,
CSV generation, provenance export, QA export — plus an end-to-end run."""

from __future__ import annotations

import json
import os

import pytest
from openpyxl import load_workbook

from abcc_recovery import cli
from abcc_recovery.export import csv as csv_export
from abcc_recovery.export import excel as excel_export
from abcc_recovery.processing.lineage import build_provenance
from abcc_recovery.processing.normalize import normalize_records
from abcc_recovery.qa import reconcile as reconcile_mod
from abcc_recovery.qa import validate as validate_mod
from abcc_recovery.recovery.downloader import Downloader
from abcc_recovery.recovery.source_discovery import SourceDiscovery
from abcc_recovery.recovery.wayback import WaybackClient
from tests.conftest import FakeResponse, FakeSession, cdx_payload
from tests.fixtures import make_xlsx

PDF_URL = (
    "https://www.abcc.gov.au/sites/default/files/abcc_agreement_clauses_spreadsheet.pdf"
)


def _dataset():
    raw = [
        {
            "raw_id": f"RAW-{i:06d}",
            "source_id": "SRC-0001",
            "source_file": "fixture.xlsx",
            "sheet": "Clauses",
            "page": None,
            "row_number": i + 1,
            "seq": i,
            "fields": {"Clause No": no, "Clause Text": text},
        }
        for i, (no, text) in enumerate(
            [
                ("1", "First clause text for export testing."),
                ("2", "Second clause text for export testing."),
                ("2", "Second clause text for export testing."),
            ],
            start=1,
        )
    ]
    clean = normalize_records(raw)
    findings = validate_mod.validate_records(clean, raw)
    recon = reconcile_mod.reconcile(
        raw,
        clean,
        source_record_count=3,
        exact_duplicate_groups=findings["exact_duplicate_groups"],
    )
    provenance = build_provenance(
        "SRC-0001",
        {
            "original_url": PDF_URL,
            "archive_url": "https://web.archive.org/web/20150101000000id_/" + PDF_URL,
            "archive_timestamp": "20150101000000",
            "saved_path": "/data/source/abcc_agreement_clauses_spreadsheet.pdf",
            "detected_file_type": "xlsx",
            "file_size": 1234,
            "sha256": "ab" * 32,
        },
        retrieval_date="2024-01-01T00:00:00+00:00",
        extraction_method="openpyxl",
    )
    readme = {
        "purpose": "test",
        "source": PDF_URL,
        "recovery_method": "wayback",
        "recovery_date": "2024-01-01",
        "extraction_method": "openpyxl",
        "limitations": "test limitations",
    }
    return raw, clean, findings, recon, [provenance.export_row()], readme


def test_workbook_creation_and_sheets(tmp_path):
    raw, clean, findings, recon, prov, readme = _dataset()
    out = str(tmp_path / excel_export.WORKBOOK_NAME)
    excel_export.create_workbook(raw, clean, findings, recon, prov, readme, out)
    assert os.path.exists(out)

    wb = load_workbook(out)
    assert wb.sheetnames == [
        "01_README",
        "02_RAW_DATA",
        "03_CLEAN_DATA",
        "04_QA_REPORT",
        "05_SOURCE_PROVENANCE",
    ]
    # XLS-006: nothing hidden
    assert all(ws.sheet_state == "visible" for ws in wb.worksheets)
    # XLS-001 / XLS-002 on data sheets
    for name in ("02_RAW_DATA", "03_CLEAN_DATA"):
        ws = wb[name]
        assert ws.auto_filter.ref is not None
        assert ws.freeze_panes == "A2"
    # README contains the section-24 legal disclaimer
    readme_text = "\n".join(
        str(c.value) for row in wb["01_README"].iter_rows() for c in row if c.value
    )
    assert "does not constitute legal advice" in readme_text
    assert "historical" in readme_text.lower()
    # 02_RAW_DATA has all raw rows + header
    assert wb["02_RAW_DATA"].max_row == len(raw) + 1
    assert wb["03_CLEAN_DATA"].max_row == len(clean) + 1


def test_provenance_export(tmp_path):
    raw, clean, findings, recon, prov, readme = _dataset()
    out = str(tmp_path / "wb.xlsx")
    excel_export.create_workbook(raw, clean, findings, recon, prov, readme, out)
    ws = load_workbook(out)["05_SOURCE_PROVENANCE"]
    headers = [c.value for c in ws[1]]
    assert headers[:2] == ["source_id", "original_url"]
    values = [c.value for c in ws[2]]
    row = dict(zip(headers, values))
    assert row["source_id"] == "SRC-0001"
    assert row["original_url"] == PDF_URL
    assert row["sha256"] == "ab" * 32
    assert row["extraction_method"] == "openpyxl"


def test_qa_export_contains_counts_and_flagged_records(tmp_path):
    raw, clean, findings, recon, prov, readme = _dataset()
    out = str(tmp_path / "wb.xlsx")
    excel_export.create_workbook(raw, clean, findings, recon, prov, readme, out)
    ws = load_workbook(out)["04_QA_REPORT"]
    text = "\n".join(
        str(c.value) for row in ws.iter_rows() for c in row if c.value is not None
    )
    assert "Total records" in text
    assert "Extracted records: 3" in text  # REC-004 reconciliation summary
    assert "Manual review required" in text
    assert "REC-000003" in text  # flagged duplicate listed with reason
    assert "duplicate" in text


def test_csv_generation(tmp_path):
    _raw, clean, _f, _r, _p, _m = _dataset()
    out = str(tmp_path / csv_export.CSV_NAME)
    csv_export.write_csv(clean, out)
    with open(out, encoding="utf-8") as fh:
        lines = fh.read().splitlines()
    assert lines[0].startswith("record_id,raw_id,source_id")
    assert len(lines) == len(clean) + 1
    assert "REC-000001" in lines[1]


def test_csv_column_order_matches_clean_sheet(tmp_path):
    """CSV column order must match 03_CLEAN_DATA (qa_status last)."""
    raw, clean, findings, recon, prov, readme = _dataset()
    xlsx_out = str(tmp_path / excel_export.WORKBOOK_NAME)
    csv_out = str(tmp_path / csv_export.CSV_NAME)
    excel_export.create_workbook(raw, clean, findings, recon, prov, readme, xlsx_out)
    csv_export.write_csv(clean, csv_out)
    sheet_headers = [c.value for c in load_workbook(xlsx_out)["03_CLEAN_DATA"][1]]
    with open(csv_out, encoding="utf-8") as fh:
        csv_headers = fh.readline().strip().split(",")
    assert csv_headers == sheet_headers
    assert csv_headers[-1] == "qa_status"


def test_export_failure_exit_code(tmp_path):
    """ERR-005: export without data must fail with a non-zero code."""
    ctx = cli.PipelineContext(str(tmp_path))
    assert cli.stage_export(ctx) == cli.EXIT_EXPORT_FAILED


# ----------------------------------------------------------------------
# End-to-end run with a fully mocked archive (offline)
# ----------------------------------------------------------------------
def test_full_pipeline_run_success(config_path, tmp_path):
    xlsx_bytes = open(make_xlsx(str(tmp_path / "fixture_src.xlsx")), "rb").read()
    download_session = FakeSession(
        {
            "web.archive.org/web": FakeResponse(200, content=xlsx_bytes),
        }
    )
    wayback_session = FakeSession(
        {
            "cdx": FakeResponse(
                json_data=cdx_payload(
                    [
                        [
                            "au,gov,abcc)/pdf",
                            "20160101000000",
                            PDF_URL,
                            "application/pdf",
                            "200",
                            "D1",
                            "10",
                        ],
                    ]
                )
            ),
        }
    )
    discovery = SourceDiscovery(
        WaybackClient(session=wayback_session, sleep=lambda s: None),
        Downloader(session=download_session),
    )
    ctx = cli.PipelineContext(str(tmp_path), config_path)
    code = cli.stage_run(ctx, discovery=discovery)
    assert code == cli.EXIT_OK

    xlsx_out = tmp_path / "data" / "output" / excel_export.WORKBOOK_NAME
    csv_out = tmp_path / "data" / "output" / csv_export.CSV_NAME
    assert xlsx_out.exists() and csv_out.exists()
    assert load_workbook(str(xlsx_out)).sheetnames == [
        "01_README",
        "02_RAW_DATA",
        "03_CLEAN_DATA",
        "04_QA_REPORT",
        "05_SOURCE_PROVENANCE",
    ]

    recovery = json.loads((tmp_path / "reports" / "recovery_report.json").read_text())
    assert recovery["status"] == "RECOVERED"
    assert recovery["selected_source"]["sha256"]
    qa = json.loads((tmp_path / "reports" / "qa_report.json").read_text())
    assert qa["clean_record_count"] == 7
    assert qa["reconciliation"]["summary_lines"]
