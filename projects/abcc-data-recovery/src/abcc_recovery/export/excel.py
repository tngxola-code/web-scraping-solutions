"""Excel workbook export (SPEC sections 13, 14 and 24).

Creates ``ABCC_Agreement_Clauses.xlsx`` with EXACTLY five worksheets:
01_README, 02_RAW_DATA, 03_CLEAN_DATA, 04_QA_REPORT, 05_SOURCE_PROVENANCE.
"""

from __future__ import annotations

import logging
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

log = logging.getLogger("abcc_recovery.export.excel")

WORKBOOK_NAME = "ABCC_Agreement_Clauses.xlsx"
SHEET_README = "01_README"
SHEET_RAW = "02_RAW_DATA"
SHEET_CLEAN = "03_CLEAN_DATA"
SHEET_QA = "04_QA_REPORT"
SHEET_PROVENANCE = "05_SOURCE_PROVENANCE"

LEGAL_DISCLAIMER = (
    "This dataset contains historical material recovered from former "
    "Australian Building and Construction Commission resources. It is "
    "provided for historical/data-recovery purposes and does not constitute "
    "legal advice or confirmation of current regulatory requirements."
)

MAX_COLUMN_WIDTH = 80


class ExportError(Exception):
    """ERR-005: a failed export must produce a non-zero process exit code."""


def _field_columns(records: list[dict[str, Any]]) -> list[str]:
    """Union of field names, in first-seen order (XLS-005 consistent headers)."""
    columns: list[str] = []
    for rec in records:
        for key in rec.get("fields") or {}:
            if key not in columns:
                columns.append(key)
    return columns


def _write_table(
    ws: Any,
    meta_columns: list[str],
    records: list[dict[str, Any]],
    extra_columns: Optional[list[str]] = None,
) -> None:
    fields = _field_columns(records)
    headers = meta_columns + fields + (extra_columns or [])
    ws.append(headers)
    for rec in records:
        row = [rec.get(c) for c in meta_columns]
        row += [(rec.get("fields") or {}).get(f) for f in fields]
        row += [rec.get(c) for c in (extra_columns or [])]
        ws.append(row)
    _format_sheet(ws, headers)


def _format_sheet(ws: Any, headers: list[str]) -> None:
    """XLS-001 filters, XLS-002 frozen header, XLS-003 wrap, XLS-004 widths."""
    header_font = Font(bold=True)
    wrap = Alignment(wrap_text=True, vertical="top")
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        letter = get_column_letter(col_idx)
        width = max(len(str(header)) + 2, 12)
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            value = row[0].value
            if value is not None:
                width = max(width, min(len(str(value)) + 2, MAX_COLUMN_WIDTH))
            row[0].alignment = wrap  # XLS-003
        ws.column_dimensions[letter].width = width  # XLS-004
    last_col = get_column_letter(max(1, len(headers)))
    ws.auto_filter.ref = f"A1:{last_col}{max(1, ws.max_row)}"  # XLS-001
    ws.freeze_panes = "A2"  # XLS-002


def _write_readme(ws: Any, info: dict[str, Any]) -> None:
    """01_README content (SPEC section 13 + section 24 disclaimer)."""
    ws.append(["ABCC Agreement Clauses — Recovered Dataset"])
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    rows = [
        ("Purpose", info.get("purpose", "")),
        ("Source", info.get("source", "")),
        ("Recovery method", info.get("recovery_method", "")),
        ("Recovery date", info.get("recovery_date", "")),
        ("Extraction method", info.get("extraction_method", "")),
        ("Important limitations", info.get("limitations", "")),
        ("Legal disclaimer", LEGAL_DISCLAIMER),
    ]
    for label, value in rows:
        ws.append([label, value])
        ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
        ws.cell(row=ws.max_row, column=2).alignment = Alignment(
            wrap_text=True, vertical="top"
        )
    ws.column_dimensions["A"].width = 24
    ws.column_dimensions["B"].width = 100


def _write_qa(
    ws: Any, qa_findings: dict[str, Any], reconciliation: dict[str, Any]
) -> None:
    """04_QA_REPORT content (SPEC section 13)."""
    ws.append(["Metric", "Value"])
    ws.cell(row=1, column=1).font = Font(bold=True)
    ws.cell(row=1, column=2).font = Font(bold=True)
    rows: list[tuple[str, Any]] = [
        ("Total records", qa_findings.get("total_records", 0)),
        ("Exact duplicate count", qa_findings.get("exact_duplicate_count", 0)),
        ("Near duplicate pairs", len(qa_findings.get("near_duplicates", []))),
    ]
    for field_name, count in sorted(
        (qa_findings.get("missing_value_stats") or {}).items()
    ):
        rows.append((f"Missing values: {field_name}", count))
    for status, count in sorted((qa_findings.get("status_counts") or {}).items()):
        rows.append((f"Status: {status}", count))
    rows.append(
        (
            "Records requiring manual review",
            reconciliation.get("manual_review_required", 0),
        )
    )
    # REC-004 reconciliation summary
    for line in reconciliation.get("summary_lines", []):
        rows.append(("Reconciliation", line))
    rows.append(("Reconciled", reconciliation.get("reconciled")))
    for label, value in rows:
        ws.append([label, value])
    # flagged records with human-readable reasons (QA-007)
    ws.append([])
    ws.append(["Flagged record", "Review reason"])
    ws.cell(row=ws.max_row, column=1).font = Font(bold=True)
    ws.cell(row=ws.max_row, column=2).font = Font(bold=True)
    for item in qa_findings.get("manual_review_records", []):
        ws.append([item["record_id"], "; ".join(item["review_notes"])])
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 100
    ws.freeze_panes = "A2"


def _write_provenance(ws: Any, provenance_rows: list[dict[str, Any]]) -> None:
    """05_SOURCE_PROVENANCE content (SPEC section 13 / LIN-004)."""
    headers = [
        "source_id",
        "original_url",
        "recovered_url",
        "archive_timestamp",
        "filename",
        "file_type",
        "file_size",
        "sha256",
        "retrieval_date",
        "extraction_method",
    ]
    ws.append(headers)
    for row in provenance_rows:
        ws.append([row.get(h) for h in headers])
    _format_sheet(ws, headers)


def create_workbook(
    raw_records: list[dict[str, Any]],
    clean_records: list[dict[str, Any]],
    qa_findings: dict[str, Any],
    reconciliation: dict[str, Any],
    provenance_rows: list[dict[str, Any]],
    readme_info: dict[str, Any],
    out_path: str,
) -> str:
    """Build and save the five-sheet workbook; raises ExportError on failure."""
    try:
        wb = Workbook()
        ws_readme = wb.active
        ws_readme.title = SHEET_README
        _write_readme(ws_readme, readme_info)

        ws_raw = wb.create_sheet(SHEET_RAW)
        _write_table(
            ws_raw,
            [
                "raw_id",
                "source_id",
                "source_file",
                "sheet",
                "page",
                "row_number",
                "seq",
            ],
            raw_records,
        )

        ws_clean = wb.create_sheet(SHEET_CLEAN)
        _write_table(
            ws_clean,
            [
                "record_id",
                "raw_id",
                "source_id",
                "source_file",
                "sheet",
                "page",
                "row_number",
            ],
            clean_records,
            extra_columns=["qa_status"],
        )

        ws_qa = wb.create_sheet(SHEET_QA)
        _write_qa(ws_qa, qa_findings, reconciliation)

        ws_prov = wb.create_sheet(SHEET_PROVENANCE)
        _write_provenance(ws_prov, provenance_rows)

        # XLS-006: no hidden sheets
        for ws in wb.worksheets:
            ws.sheet_state = "visible"

        wb.save(out_path)
    except Exception as exc:  # noqa: BLE001 - ERR-005
        raise ExportError(f"workbook export failed: {exc}") from exc
    log.info("[export] Workbook created: %s", out_path)
    return out_path
